"""
generar_informe.py - Informe final (informe.docx) y presentacion (presentacion.pptx)
Persona 4 (escenarios, datos, decision e integracion)

Lee las tablas y figuras de resultados/ (correr antes escenarios.py,
escenario_e.py, montecarlo.py, decision.py y main.py) y arma el informe con
las normas de la universidad:
  - Times New Roman 12, interlineado 1.5, texto justificado, todo en negro.
  - Titulos numerados 1, 1.1, 1.1.1.
  - Tablas solo con lineas horizontales; titulo arriba (negrita, centrado)
    y fuente abajo (centrada). Igual para figuras.
  - Diagramas en blanco y negro; solo las graficas llevan color.
Al final comprueba que cada archivo de resultados/ este citado en el texto.

Uso:  python informe/generar_informe.py
"""

import os
import sys

import pandas as pd

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

import comun  # noqa: E402  (configura matplotlib en modo Agg)
import params as P  # noqa: E402
import modelo_euler as me  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402,F401
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

CARPETA = os.path.join(RAIZ, "informe")
FIG_INF = os.path.join(CARPETA, "figuras")
FIG = P.FIGURAS
TAB = P.TABLAS
NEGRO = RGBColor(0, 0, 0)
FUENTE_PROPIA = "Fuente: Elaboración propia 2026"
FUENTE_SINT = "Fuente: Elaboración propia 2026 en base a datos sintéticos (datos/README.md)"


def tabla_csv(nombre):
    return pd.read_csv(os.path.join(TAB, nombre))


def n(x, d=2):
    """Numero con coma decimal."""
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, d=0):
    return n(100 * x, d) + " %"


# ---------------------------------------------------------------------------
# Diagramas en blanco y negro
# ---------------------------------------------------------------------------
def _caja(ax, x, y, texto, w=2.2, h=0.9):
    from matplotlib.patches import FancyBboxPatch
    ax.add_patch(FancyBboxPatch((x - w / 2, y - h / 2), w, h, boxstyle="round,pad=0.05",
                                fc="white", ec="black", lw=1.2))
    ax.text(x, y, texto, ha="center", va="center", fontsize=9, color="black")


def _flecha(ax, a, b):
    ax.annotate("", xy=b, xytext=a, arrowprops=dict(arrowstyle="->", color="black", lw=1.2))


def diagrama_pipeline():
    plt = comun.plt
    fig, ax = plt.subplots(figsize=(11, 1.9))
    pasos = [("Contexto\nhora, clima,\nevento, tráfico", 0.9),
             ("Red bayesiana\nP(escasez), D0, S0\n(Persona 2)", 3.4),
             ("Modelo de Euler\nprecio, demanda,\noferta (Persona 1)", 5.9),
             ("Filtro de Kalman\nprecio estimado\n± 2 desv. (Persona 3)", 8.4),
             ("Decisión\nMANTENER / SUBIR /\nINCENTIVAR (Persona 4)", 10.9)]
    for texto, x in pasos:
        _caja(ax, x, 1, texto, w=2.2, h=1.2)
    for (_, x1), (_, x2) in zip(pasos, pasos[1:]):
        _flecha(ax, (x1 + 1.12, 1), (x2 - 1.12, 1))
    ax.annotate("", xy=(10.9, 0.38), xytext=(3.4, 0.38),
                arrowprops=dict(arrowstyle="->", color="black", lw=1, ls="--",
                                connectionstyle="arc3,rad=0.12"))
    ax.text(7.15, 0.02, "P(escasez) va directo a la decisión", ha="center", fontsize=8)
    ax.set_xlim(-0.4, 12.2)
    ax.set_ylim(-0.1, 1.8)
    ax.axis("off")
    ruta = os.path.join(FIG_INF, "diagrama_pipeline.png")
    fig.savefig(ruta, dpi=200, bbox_inches="tight")
    plt.close(fig)
    return ruta


def diagrama_dag():
    plt = comun.plt
    pos = {"Hora": (1, 3), "Clima": (3, 3), "Evento": (5, 3), "Costos": (7, 3),
           "Trafico": (2, 2), "Demanda": (4, 1.6), "Oferta": (6, 1.6), "Escasez": (5, 0.4)}
    aristas = [("Hora", "Trafico"), ("Clima", "Trafico"), ("Hora", "Demanda"),
               ("Clima", "Demanda"), ("Evento", "Demanda"), ("Clima", "Oferta"),
               ("Trafico", "Oferta"), ("Costos", "Oferta"), ("Demanda", "Escasez"),
               ("Oferta", "Escasez")]
    fig, ax = plt.subplots(figsize=(8, 4))
    for a, b in aristas:
        ax.annotate("", xy=pos[b], xytext=pos[a],
                    arrowprops=dict(arrowstyle="->", color="black", lw=1, shrinkA=18, shrinkB=18))
    for nombre, (x, y) in pos.items():
        ax.text(x, y, nombre.replace("Trafico", "Tráfico"), ha="center", va="center", fontsize=10,
                bbox=dict(boxstyle="round,pad=0.4", fc="white", ec="black"))
    ax.set_xlim(0.2, 7.8)
    ax.set_ylim(0, 3.5)
    ax.axis("off")
    ruta = os.path.join(FIG_INF, "diagrama_red_bayesiana.png")
    fig.savefig(ruta, dpi=200, bbox_inches="tight")
    plt.close(fig)
    return ruta


# ---------------------------------------------------------------------------
# Formato del documento
# ---------------------------------------------------------------------------
def _fuente_run(run, negrita=None, cursiva=None, tam=12):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(tam)
    run.font.color.rgb = NEGRO
    if negrita is not None:
        run.bold = negrita
    if cursiva is not None:
        run.italic = cursiva


def configurar_estilos(doc):
    for nombre in ["Normal", "Heading 1", "Heading 2", "Heading 3", "Title", "List Bullet",
                   "List Number", "Caption"]:
        st = doc.styles[nombre]
        st.font.name = "Times New Roman"
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(att), "Times New Roman")
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(att)) is not None:
                del rfonts.attrib[qn(att)]
        st.font.size = Pt(12)
        st.font.color.rgb = NEGRO
        st.font.italic = False
        pf = st.paragraph_format
        pf.line_spacing = 1.5
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_after = Pt(6)
        if nombre.startswith("Heading") or nombre == "Title":
            st.font.bold = True
            pf.space_before = Pt(12)
            pf.keep_with_next = True
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(3)
        s.right_margin = Cm(2.5)


def numero_de_pagina(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)
    _fuente_run(run)


class Informe:
    def __init__(self):
        self.doc = Document()
        configurar_estilos(self.doc)
        numero_de_pagina(self.doc)
        self.n_tabla = 0
        self.n_figura = 0
        self.texto_total = []

    # --- texto -----------------------------------------------------------
    def titulo(self, numero, texto, nivel=1):
        h = self.doc.add_heading(level=nivel)
        h.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _fuente_run(h.add_run(f"{numero} {texto}" if numero else texto), negrita=True)
        self.texto_total.append(texto)

    def parrafo(self, texto, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY, negrita=False, cursiva=False):
        p = self.doc.add_paragraph()
        p.alignment = alineacion
        # **negrita** simple dentro del texto
        partes = texto.split("**")
        for i, parte in enumerate(partes):
            if parte:
                _fuente_run(p.add_run(parte), negrita=negrita or (i % 2 == 1), cursiva=cursiva)
        self.texto_total.append(texto)
        return p

    def vinetas(self, items, estilo="List Bullet"):
        for it in items:
            p = self.doc.add_paragraph(style=estilo)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            partes = it.split("**")
            for i, parte in enumerate(partes):
                if parte:
                    _fuente_run(p.add_run(parte), negrita=(i % 2 == 1))
            self.texto_total.append(it)

    def _leyenda(self, texto, negrita):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = negrita
        p.paragraph_format.space_after = Pt(2 if negrita else 10)
        _fuente_run(p.add_run(texto), negrita=negrita)
        self.texto_total.append(texto)

    # --- tablas y figuras -------------------------------------------------
    def tabla(self, titulo, encabezados, filas, fuente=FUENTE_PROPIA, anchos=None):
        self.n_tabla += 1
        self._leyenda(f"Tabla {self.n_tabla}. {titulo}", True)
        t = self.doc.add_table(rows=1, cols=len(encabezados))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, e in enumerate(encabezados):
            self._celda(t.rows[0].cells[j], e, True)
        for fila in filas:
            celdas = t.add_row().cells
            for j, v in enumerate(fila):
                self._celda(celdas[j], v, False)
        if anchos:
            for fila in t.rows:
                for j, w in enumerate(anchos):
                    fila.cells[j].width = Cm(w)
        self._bordes_horizontales(t)
        self._leyenda(fuente, False)
        return self.n_tabla

    def _celda(self, celda, texto, negrita):
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        _fuente_run(p.add_run(str(texto)), negrita=negrita)
        self.texto_total.append(str(texto))

    @staticmethod
    def _bordes_horizontales(t):
        tblPr = t._tbl.tblPr
        bordes = OxmlElement("w:tblBorders")
        for lado, valor in (("top", "single"), ("left", "nil"), ("bottom", "single"),
                            ("right", "nil"), ("insideH", "single"), ("insideV", "nil")):
            el = OxmlElement(f"w:{lado}")
            el.set(qn("w:val"), valor)
            if valor == "single":
                el.set(qn("w:sz"), "6")
                el.set(qn("w:color"), "000000")
            bordes.append(el)
        tblPr.append(bordes)

    def figura(self, titulo, ruta, fuente=FUENTE_PROPIA, ancho=15.5):
        self.n_figura += 1
        self._leyenda(f"Figura {self.n_figura}. {titulo}", True)
        self.doc.add_picture(ruta, width=Cm(ancho))
        p = self.doc.paragraphs[-1]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        self._leyenda(fuente, False)
        return self.n_figura

    def salto(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Contenido
# ---------------------------------------------------------------------------
def construir(inf):
    met = tabla_csv("metricas_escenarios.csv").set_index("escenario")
    vi = tabla_csv("kalman_valor_informacion.csv").set_index("metrica")
    marg = tabla_csv("E_frecuencias_marginales.csv")
    conj = tabla_csv("E_frecuencias_contexto.csv")
    ident = tabla_csv("E_identificabilidad.csv")
    cal = tabla_csv("E_calibracion_traduccion.csv")
    sep = tabla_csv("E_separacion_discretizacion.csv").set_index("Hora")
    redes = tabla_csv("E_red_experta_vs_calibrada.csv")
    kal = tabla_csv("E_kalman_innovaciones.csv")
    mc = tabla_csv("montecarlo_resumen.csv")
    dec = tabla_csv("decision_resultado_elegido.csv").set_index("escenario")
    eX = tabla_csv("decision_eleccion_X.csv").set_index("X")
    eY = tabla_csv("decision_eleccion_Y.csv").set_index("Y")
    perm = tabla_csv("decision_permanencia.csv")
    est = me.estabilidad(P.PARAMS)
    X, Y = P.UMBRAL_X, P.UMBRAL_Y

    def kal_fila(entradas, medic, fq):
        r = kal[(kal["entradas_D0_S0"] == entradas) & (kal["mediciones"] == medic)
                & (kal["factor_q"] == fq)]
        return r.iloc[0]

    viernes_exp = redes[(redes.hora_punta == 1) & (redes.lluvia == 1) & (redes.evento == 1)].iloc[0]

    # ------------------------------------------------------------ portada
    d = inf.doc
    for _ in range(4):
        d.add_paragraph()
    inf.parrafo("PRECIO DINÁMICO DE TRANSPORTE SEGÚN OFERTA Y DEMANDA", WD_ALIGN_PARAGRAPH.CENTER,
                negrita=True)
    inf.parrafo("Modelo de Euler, red bayesiana y filtro de Kalman integrados en una regla de "
                "decisión de tarifa", WD_ALIGN_PARAGRAPH.CENTER)
    d.add_paragraph()
    inf.parrafo("Informe final", WD_ALIGN_PARAGRAPH.CENTER, negrita=True)
    inf.parrafo("Asignatura: Modelado", WD_ALIGN_PARAGRAPH.CENTER)
    d.add_paragraph()
    for linea in ["Persona 1, modelo dinámico (Euler): [completar nombre]",
                  "Persona 2, red bayesiana: [completar nombre]",
                  "Persona 3, filtro de Kalman: [completar nombre]",
                  "Persona 4, escenarios, datos, decisión e integración: Rubén Cordero"]:
        inf.parrafo(linea, WD_ALIGN_PARAGRAPH.CENTER)
    d.add_paragraph()
    inf.parrafo("Cochabamba, 2026", WD_ALIGN_PARAGRAPH.CENTER)
    inf.salto()

    # ------------------------------------------------------------ 1
    inf.titulo("1.", "Introducción")
    inf.parrafo(
        "Las plataformas de transporte ajustan el precio de un viaje según la relación entre "
        "pasajeros que piden un vehículo (demanda) y conductores disponibles (oferta). Cuando "
        "llueve, en hora punta o cuando hay un evento masivo, la demanda sube o la oferta baja, y "
        "el precio reacciona. Este trabajo construye un sistema que reproduce ese mecanismo y lo "
        "usa para decidir, en cada intervalo de cinco minutos, si conviene mantener el precio, "
        "subirlo o pagar un incentivo a los conductores.")
    inf.parrafo(
        "El sistema tiene tres capas, cada una a cargo de un integrante: un modelo dinámico "
        "resuelto con el método de Euler, que calcula cómo evoluciona el precio; una red "
        "bayesiana, que traduce el contexto (hora, clima, eventos, tráfico y costos) en "
        "probabilidades de demanda, oferta y escasez; y un filtro de Kalman, que estima el precio "
        "y la demanda reales a partir de mediciones con ruido. La cuarta parte, presentada aquí, "
        "integra las tres capas, corre los escenarios A a E del caso, calibra la red con datos y "
        "define la regla de decisión que responde la pregunta 7.")
    inf.parrafo(
        "Los datos del Escenario E son sintéticos: se generaron con el propio modelo porque no "
        "se encontraron series públicas de precios de plataformas en Cochabamba. Esto se indica "
        "en cada tabla y figura que los usa y se discute en la sección 4.")

    # ------------------------------------------------------------ 2
    inf.titulo("2.", "Descripción del sistema")
    inf.parrafo(
        "La Figura 1 resume cómo se conectan las capas. Cada capa vive en su propio archivo y "
        "la integración (sección 5) solo las importa, sin modificar su lógica.")
    inf.figura("Flujo del sistema integrado", diagrama_pipeline())

    inf.titulo("2.1", "Modelo dinámico con el método de Euler (Persona 1)", 2)
    inf.parrafo(
        "Archivos: modelo_euler.py y precio_dinamico_transporte.ipynb. En cada intervalo k "
        "la demanda y la oferta responden al precio y el precio se corrige según el exceso de "
        "demanda:")
    inf.parrafo("D_k = D0 + β (p_ref − p_k) + ε_k;   S_k = S0 + γ (p_k − p_costo);   "
                "p_(k+1) = max(p_costo, p_k + Δt α (D_k − S_k))", WD_ALIGN_PARAGRAPH.CENTER)
    inf.parrafo(
        f"con α = {n(P.PARAMS['alpha'], 1)}, β = {n(P.PARAMS['beta'], 0)}, "
        f"γ = {n(P.PARAMS['gamma'], 1)}, p_ref = {n(P.PARAMS['p_ref'], 0)} Bs, "
        f"p_costo = {n(P.PARAMS['p_costo'], 0)} Bs, Δt = {n(P.PARAMS['dt'], 0)} min, "
        f"T = {n(P.PARAMS['T'], 0)} min y p0 = {n(P.PARAMS['p0'], 0)} Bs. Igualando demanda y oferta "
        "se obtiene el precio de equilibrio p* = (D0 − S0 + β p_ref + γ p_costo) / (β + γ), que con "
        f"D0 = S0 = 50 vale {n(met.loc['A', 'p_eq_teorico_base'])} Bs. El error respecto del "
        f"equilibrio se multiplica en cada paso por r = 1 − Δt α (β + γ) = {n(est['r'])}; como "
        f"Δt α (β + γ) = {n(est['producto'])} < 2 el sistema es estable, y como r < 0 el precio "
        "oscila alrededor del equilibrio. La función simular() acepta D0 y S0 como números, "
        "arreglos o funciones del paso k, lo que permite los escenarios con hora punta o lluvia "
        "por tramos. El cuaderno de la Persona 1 incluye además el análisis de sensibilidad de α, "
        "β y γ (pregunta 1) y la matriz de transición F que usa el filtro de Kalman.")

    inf.titulo("2.2", "Red bayesiana de factores contextuales (Persona 2)", 2)
    inf.parrafo(
        "Archivos: Bayes/red.py, Bayes/consultas.py, Bayes/traduccion.py y Bayes/dibujar.py. "
        "La red (Figura 2) tiene cuatro nodos de contexto (Hora, Clima, Evento y Costos), un "
        "nodo intermedio (Tráfico), dos nodos de mercado (Demanda y Oferta, con niveles baja, "
        "media y alta) y el nodo Escasez. Las tablas de probabilidad son supuestos expertos "
        "justificados en red.py y la inferencia se programó desde cero (multiplicar y "
        "normalizar); pgmpy se usa solo para verificar. Para la consulta del documento (viernes "
        "19:00, lluvia y evento masivo) la red entrega P(escasez) = "
        f"{n(viernes_exp['P_escasez_red_experta'], 3)}. traduccion.py convierte las categorías "
        "de la red en los valores D0 y S0 que necesita el modelo de Euler.")
    inf.figura("Red bayesiana: nodos y aristas", diagrama_dag(),
               fuente="Fuente: Elaboración propia 2026 en base a Bayes/red.py", ancho=12)

    inf.titulo("2.3", "Filtro de Kalman (Persona 3)", 2)
    inf.parrafo(
        "Archivo: filtro_kalman.py. El estado es x_k = (p_k, D_k). La transición x_k = F x_(k−1) "
        "+ B u_k + w_k usa la matriz F derivada del modelo de Euler, y la observación y_k = H x_k "
        "+ v_k mide solo el precio (H = [1, 0]) o el precio y las solicitudes (H = I2). El filtro "
        "está escrito a mano (predicción, ganancia y corrección) e incluye los experimentos de "
        "calibración de Q y R (pregunta 3), valor de la información (pregunta 5) y valores "
        "atípicos.")

    inf.titulo("2.4", "Integración (Persona 4)", 2)
    inf.parrafo(
        "Archivos: params.py (única fuente de parámetros y definición de los escenarios), "
        "comun.py, escenarios.py, escenario_e.py, montecarlo.py, decision.py, main.py y este "
        "informe (informe/generar_informe.py). La única modificación al código de otra capa son "
        "las constantes VALORES_D0 y VALORES_S0 de Bayes/traduccion.py, calibradas con los datos "
        "(sección 3.4.3).")

    # ------------------------------------------------------------ 3
    inf.titulo("3.", "Escenarios y resultados")
    inf.parrafo(
        "Los escenarios A, B y C se simulan durante 300 minutos con el modelo de Euler. B y C "
        "aplican la perturbación solo entre los minutos 60 y 180, de modo que se ve tanto la "
        "reacción como la recuperación (Figura 4). La Tabla 1 reúne las métricas "
        "(resultados/tablas/metricas_escenarios.csv, generada por escenarios.py).")
    filas = []
    for e in "ABC":
        r = met.loc[e]
        filas.append([e, n(r["p_eq_teorico_base"]), n(r["p_eq_teorico_ventana"]),
                      n(r["precio_medio"]), n(r["precio_max"]),
                      n(r["tiempo_hasta_equilibrio_min"], 0),
                      "-" if pd.isna(r["tiempo_equilibrio_ventana_min"])
                      else n(r["tiempo_equilibrio_ventana_min"], 0),
                      int(r["intervalos_con_escasez"]), n(r["servicios_realizados"], 0)])
    inf.tabla("Métricas de los escenarios A, B y C",
              ["Esc.", "p* base (Bs)", "p* ventana (Bs)", "Precio medio", "Precio máx.",
               "t. equil. (min)", "t. equil. ventana", "Interv. escasez", "Servicios"], filas)
    inf.parrafo(
        "p* base es el equilibrio con D0 = S0 = 50 y p* ventana el equilibrio durante la "
        "perturbación. El tiempo hasta el equilibrio es el primer instante desde el cual el "
        "precio no sale de una banda de ±1 % alrededor de p*; en B y C se mide hasta el "
        "equilibrio final (después de la ventana) y, en la columna siguiente, desde el minuto 60 "
        "hasta el equilibrio de la ventana. Se cuenta escasez cuando faltan al menos un vehículo "
        "(D − S ≥ 1).")

    inf.titulo("3.1", "Escenario A: parámetros nominales", 2)
    inf.parrafo(
        f"Con p0 = 10 Bs el precio no arranca en equilibrio: p* = {n(met.loc['A', 'p_eq_teorico_base'])} "
        f"Bs. Como r = {n(est['r'])}, el precio cruza el equilibrio en cada paso con oscilaciones "
        f"que se reducen un 25 % por intervalo, y entra en la banda de ±1 % a los "
        f"{n(met.loc['A', 'tiempo_hasta_equilibrio_min'], 0)} minutos (Figura 3). Durante ese "
        f"transitorio hay {int(met.loc['A', 'intervalos_con_escasez'])} intervalos con escasez. "
        "No es correcto describir el Escenario A como un precio que permanece estable: converge "
        "oscilando.")
    inf.figura("Escenario A: convergencia oscilante hacia p*",
               os.path.join(FIG, "escenario_A_convergencia.png"))

    inf.titulo("3.2", "Escenarios B (hora punta) y C (lluvia)", 2)
    inf.parrafo(
        f"En B la demanda base sube a D0 = 100 y el equilibrio de la ventana pasa a "
        f"{n(met.loc['B', 'p_eq_teorico_ventana'])} Bs; el primer salto lleva el precio a "
        f"{n(met.loc['B', 'precio_max'])} Bs por la sobrerreacción propia de r < 0 y tarda "
        f"{n(met.loc['B', 'tiempo_equilibrio_ventana_min'], 0)} minutos en asentarse. En C la "
        f"oferta base cae a S0 = 30 y el equilibrio sube a {n(met.loc['C', 'p_eq_teorico_ventana'])} "
        f"Bs, con máximo de {n(met.loc['C', 'precio_max'])} Bs. Al terminar la ventana ambos "
        "vuelven al equilibrio base con una nueva oscilación (Figura 4). B realiza más servicios "
        "que A porque la oferta responde al precio alto; C realiza menos porque faltan "
        "conductores.")
    inf.figura("Precio, demanda y oferta en los escenarios A, B y C",
               os.path.join(FIG, "escenarios_ABC.png"))

    inf.titulo("3.3", "Escenario D: estimación con filtro de Kalman", 2)
    inf.parrafo(
        "Sobre la dinámica del Escenario A se agrega un choque de demanda (desviación 3) y "
        "mediciones ruidosas del precio (0,5 Bs) y de las solicitudes (3). Se usa el filtro de la "
        "Persona 3 sin cambios (ejecutar_caso_completo y experimento_valor_informacion) y se "
        f"promedian {P.KALMAN_D['realizaciones']} realizaciones (Tabla 2, "
        "resultados/tablas/kalman_valor_informacion.csv). Las estimaciones con su intervalo "
        "se exportan a resultados/tablas/estimaciones_kalman_D.csv.")
    filas = [["Error RMS del precio (Bs)", n(vi.loc["error_precio", "solo_precio"], 3),
              n(vi.loc["error_precio", "precio_y_solicitudes"], 3)],
             ["Error RMS de la demanda", n(vi.loc["error_demanda", "solo_precio"], 3),
              n(vi.loc["error_demanda", "precio_y_solicitudes"], 3)],
             ["Cobertura ±2 desv. (precio)", pct(vi.loc["cobertura_precio", "solo_precio"], 1),
              pct(vi.loc["cobertura_precio", "precio_y_solicitudes"], 1)],
             ["Cobertura ±2 desv. (demanda)", pct(vi.loc["cobertura_demanda", "solo_precio"], 1),
              pct(vi.loc["cobertura_demanda", "precio_y_solicitudes"], 1)],
             ["Tiempo de convergencia (min)", n(vi.loc["tiempo_convergencia_min", "solo_precio"], 0),
              n(vi.loc["tiempo_convergencia_min", "precio_y_solicitudes"], 0)]]
    inf.tabla("Escenario D: valor de la información", ["Métrica", "Solo precio",
                                                        "Precio + solicitudes"], filas)
    ganancia = 1 - vi.loc["error_demanda", "precio_y_solicitudes"] / vi.loc["error_demanda", "solo_precio"]
    inf.parrafo(
        f"Medir también las solicitudes reduce el error de la demanda en {pct(ganancia)} y "
        "acorta la convergencia, mientras que el error del precio cambia poco porque el precio "
        "ya se observa directamente. Los intervalos de ±2 desviaciones cubren cerca del 95 % de "
        "los valores verdaderos, lo que indica que el filtro está bien calibrado (Figura 5).")
    inf.figura("Escenario D: solo precio frente a precio y solicitudes",
               os.path.join(FIG, "escenario_D_kalman.png"))

    inf.titulo("3.4", "Escenario E: datos de Cochabamba (sintéticos)", 2)
    inf.parrafo(
        "El Escenario E usa datos/procesado/escenario_e.csv: 1.351 intervalos de cinco minutos "
        "entre el lunes 26 de octubre y el domingo 1 de noviembre de 2026, de 06:00 a 22:00, en "
        "una ruta fija. Los datos son sintéticos (sección 4). El análisis está en "
        "escenario_e.py.")

    inf.titulo("3.4.1", "Frecuencias por contexto", 3)
    filas = []
    for _, r in conj.iterrows():
        filas.append(["sí" if r.hora_punta else "no", "sí" if r.lluvia else "no",
                      "sí" if r.evento else "no", int(r.n), n(r.precio_medio),
                      pct(r.p_tarifa_alta), pct(r.p_escasez)])
    inf.tabla("Escenario E: precio, tarifa alta y escasez por contexto (datos sintéticos)",
              ["Punta", "Lluvia", "Evento", "n", "Precio medio (Bs)", "P(tarifa alta)",
               "P(escasez)"], filas, fuente=FUENTE_SINT)
    m = marg.set_index(["variable", "valor"])
    inf.parrafo(
        "La Tabla 3 (resultados/tablas/E_frecuencias_contexto.csv) y la Figura 6 "
        "(resultados/tablas/E_frecuencias_marginales.csv) muestran un hallazgo central: la "
        "escasez definida como solicitudes − vehículos ≥ 5 casi no depende del contexto "
        f"({pct(m.loc[('hora_punta', 0), 'p_escasez'])} en valle y "
        f"{pct(m.loc[('hora_punta', 1), 'p_escasez'])} en punta). La razón es el propio "
        "mecanismo: el precio sube hasta igualar demanda y oferta, y la escasez solo aparece en "
        "los transitorios. En cambio, la tarifa alta (precio ≥ 13 Bs = 1,3 p_ref) sí refleja el "
        f"contexto: {pct(m.loc[('hora_punta', 0), 'p_tarifa_alta'], 1)} en valle frente a "
        f"{pct(m.loc[('hora_punta', 1), 'p_tarifa_alta'])} en punta, y "
        f"{pct(m.loc[('lluvia', 0), 'p_tarifa_alta'])} sin lluvia frente a "
        f"{pct(m.loc[('lluvia', 1), 'p_tarifa_alta'])} con lluvia. Por eso se usa tarifa_alta "
        "como indicador observable de escasez para calibrar el nodo Escasez de la red.")
    inf.figura("Escenario E: tarifa alta y escasez por contexto (datos sintéticos)",
               os.path.join(FIG, "E_frecuencias_contexto.png"), fuente=FUENTE_SINT)

    inf.titulo("3.4.2", "Identificabilidad de los parámetros", 3)
    filas = [[r.relacion, r.parametro, n(r.valor_nominal, 2), n(r.pendiente_regresion, 3)]
             for r in ident.itertuples()]
    inf.tabla("Pendientes de una regresión simple frente a los valores nominales",
              ["Relación", "Parámetro", "Nominal", "Regresión"], filas, fuente=FUENTE_SINT)
    inf.parrafo(
        "Aunque los datos se generaron con α, β y γ conocidos, una regresión simple no los "
        "recupera (Tabla 4, resultados/tablas/E_identificabilidad.csv): la pendiente de Δt α sale "
        "atenuada y la de −β incluso cambia de signo. Hay dos causas: el ruido de medición en "
        "la variable explicativa atenúa la pendiente, y precio y demanda se determinan al mismo "
        "tiempo (simultaneidad), de modo que en los datos el precio alto coincide con demanda "
        "alta. Esto responde la pregunta 2 sobre identificabilidad: con datos observacionales de "
        "este tipo no se pueden estimar los parámetros por mínimos cuadrados ordinarios. Se "
        "mantuvieron los valores nominales del documento.")

    inf.titulo("3.4.3", "Calibración de la red bayesiana", 3)
    crudos = cal[cal["tipo"] == "crudo"].set_index("variable")
    prop = cal[cal["tipo"] == "propuesto"].set_index("variable")
    vd = cal[cal["tipo"] == "VALORES_D0"].set_index("variable")
    vs = cal[cal["tipo"] == "VALORES_S0"].set_index("variable")
    inf.parrafo(
        "Para pasar de números a las categorías de la red hacen falta cortes de discretizar(). "
        "Los terciles de las cuentas crudas (solicitudes "
        f"{n(crudos.loc['solicitudes', 'corte_1'], 0)} / {n(crudos.loc['solicitudes', 'corte_2'], 0)}, "
        "datos/README.md) separan mal los contextos, porque las solicitudes observadas ya "
        "incluyen la respuesta al precio. Se propone discretizar la demanda y la oferta a precio "
        "de referencia, D0_hat = solicitudes − β (p_ref − precio) y S0_hat = vehículos − "
        "γ (precio − p_costo), que estiman D0 y S0 (Tabla 5, "
        "resultados/tablas/E_separacion_discretizacion.csv). Sus terciles son "
        f"{n(prop.loc['D0_hat', 'corte_1'], 1)} / {n(prop.loc['D0_hat', 'corte_2'], 1)} y "
        f"{n(prop.loc['S0_hat', 'corte_1'], 1)} / {n(prop.loc['S0_hat', 'corte_2'], 1)}.")
    inf.tabla("P(Demanda = alta) según la variable discretizada",
              ["Hora", "Con D0_hat", "Con solicitudes crudas"],
              [[h, pct(sep.loc[h, "P(Demanda=alta) con D0_hat"], 1),
                pct(sep.loc[h, "P(Demanda=alta) con solicitudes crudas"], 1)]
               for h in ["valle", "punta"]], fuente=FUENTE_SINT)
    filas = [[e, n(vd.loc[e, "valor_anterior"], 0), n(vd.loc[e, "valor_calibrado"], 0),
              n(vs.loc[e, "valor_anterior"], 0), n(vs.loc[e, "valor_calibrado"], 0)]
             for e in ["baja", "media", "alta"]]
    inf.tabla("Tabla de traducción de la red al modelo de Euler",
              ["Estado", "D0 anterior", "D0 calibrado", "S0 anterior", "S0 calibrado"], filas,
              fuente=FUENTE_SINT)
    inf.parrafo(
        "Los valores calibrados (Tabla 6, resultados/tablas/E_calibracion_traduccion.csv) son el "
        "promedio de D0_hat y S0_hat dentro de cada categoría y se aplicaron como constantes en "
        "Bayes/traduccion.py, con un comentario que cita datos/README.md. Se compararon con los "
        "valores orientativos del README de datos (D0 = 50/62/75, S0 = 38/44/48): con ellos, el "
        "precio de equilibrio que predice la red difiere del precio medio observado en 3,1 Bs "
        "(error cuadrático medio ponderado), frente a 1,2 Bs con los promedios por categoría. "
        "Los valores anteriores (D0 alta = 100, S0 baja = 30) venían de los escenarios B y C del "
        "documento. Esos escenarios son choques más fuertes que cualquier contexto de los datos "
        "(D0_hat no supera 90 y S0_hat no baja de 34), por lo que B y C se siguen definiendo "
        "directamente en params.py y no a partir de la tabla de traducción.")
    filas = []
    for r in redes.itertuples():
        filas.append([("punta" if r.hora_punta else "valle") + (", lluvia" if r.lluvia else "")
                      + (", evento" if r.evento else ""), int(r.n),
                      n(r.P_escasez_red_experta, 3), n(r.P_escasez_red_calibrada, 3),
                      pct(r.freq_tarifa_alta), pct(r.freq_escasez_DS)])
    inf.tabla("P(escasez) de la red experta y de la red calibrada frente a los datos",
              ["Contexto", "n", "Red experta", "Red calibrada", "Tarifa alta", "D − S ≥ 5"],
              filas, fuente=FUENTE_SINT)
    inf.parrafo(
        "Con estos cortes y tarifa_alta como proxy de Escasez se re-estimaron las tablas de la "
        "red con estimar_desde_datos() de la Persona 2, que combina el supuesto experto (peso de "
        "20 observaciones) con los conteos. Costos no se observa en los datos y se asumió "
        "normal. La Tabla 7 (resultados/tablas/E_red_experta_vs_calibrada.csv) compara ambas "
        "redes: la calibrada separa mejor valle seco (0,04) de hora punta (0,57), pero no "
        "alcanza la frecuencia observada en los contextos extremos porque la escasez depende "
        "solo de los tres niveles de Demanda y Oferta, que se saturan.")

    inf.titulo("3.4.4", "Filtro de Kalman sobre los datos", 3)
    a = kal_fila("nominal (D0=S0=50)", "precio + solicitudes", 30.0)
    b = kal_fila("contexto red experta", "precio + solicitudes", 30.0)
    c = kal_fila("contexto red calibrada", "precio + solicitudes", 30.0)
    s = kal_fila("contexto red calibrada", "solo precio", 30.0)
    inf.parrafo(
        "En el Escenario E no hay estado verdadero, igual que con datos reales. Por eso el filtro "
        "se evalúa con el error de predicción a un paso del precio y con las innovaciones: un "
        "filtro consistente tiene una innovación cuadrática normalizada (NIS) media igual al "
        "número de mediciones. El filtro de la Persona 3 se aplicó día por día, con R = "
        "diag(0,18; 6,25) (ruido con el que se generaron los datos) y tres opciones para D0 y "
        "S0 en el término B u: constantes nominales, o los valores que entrega la red para el "
        "contexto de cada intervalo, con la red experta o la calibrada. Se barrió el factor de "
        "escala de Q (Tabla 8, resultados/tablas/E_kalman_innovaciones.csv).")
    filas = [["Nominal (D0 = S0 = 50)", "precio + solicitudes", n(a.rmse_prediccion_1paso_precio),
              n(a.NIS_medio), pct(a.cobertura_prediccion_precio_2sd)],
             ["Red experta", "precio + solicitudes", n(b.rmse_prediccion_1paso_precio),
              n(b.NIS_medio), pct(b.cobertura_prediccion_precio_2sd)],
             ["Red calibrada", "precio + solicitudes", n(c.rmse_prediccion_1paso_precio),
              n(c.NIS_medio), pct(c.cobertura_prediccion_precio_2sd)],
             ["Red calibrada", "solo precio", n(s.rmse_prediccion_1paso_precio),
              n(s.NIS_medio), pct(s.cobertura_prediccion_precio_2sd)]]
    inf.tabla("Kalman sobre el Escenario E con Q × 30: predicción a un paso",
              ["D0 y S0", "Mediciones", "Error RMS (Bs)", "NIS medio", "Cobertura ±2 desv."],
              filas, fuente=FUENTE_SINT)
    inf.parrafo(
        "Con D0 y S0 fijos el filtro supone que el precio siempre vuelve a 7,86 Bs y en hora "
        "punta se equivoca por varios bolivianos. Al alimentar el término B u con el contexto "
        "traducido por la red calibrada, el error de predicción a un paso baja a "
        f"{n(c.rmse_prediccion_1paso_precio)} Bs y el NIS medio queda en {n(c.NIS_medio)}, cerca "
        "del valor esperado 2, con Q escalado por 30 (elegido como el factor con NIS más cercano "
        "a 2). El aumento de Q compensa que el modelo no conoce los cambios graduales de "
        "contexto. Medir también las solicitudes reduce el error de predicción de "
        f"{n(s.rmse_prediccion_1paso_precio)} a {n(c.rmse_prediccion_1paso_precio)} Bs, lo que "
        "confirma en datos el valor de la información del Escenario D. La Figura 7 muestra el "
        "viernes 30 de octubre; las estimaciones de todos los días están en "
        "resultados/tablas/E_estimaciones_kalman.csv.")
    inf.figura("Escenario E: precio observado y estimado el viernes 30 de octubre "
               "(datos sintéticos)", os.path.join(FIG, "E_kalman_viernes.png"), fuente=FUENTE_SINT)

    inf.titulo("3.5", "Monte Carlo de la red bayesiana al modelo de Euler", 2)
    inf.parrafo(
        "Para llevar la incertidumbre del contexto al precio, montecarlo.py toma 500 pares (D0, "
        "S0) de P(Demanda, Oferta | contexto) con muestrear_parametros() y corre simular() con "
        "cada uno. La Figura 8 muestra la mediana y los percentiles 5 y 95 del precio "
        "(resultados/tablas/montecarlo_bandas.csv) y la Tabla 9 el resumen "
        "(resultados/tablas/montecarlo_resumen.csv).")
    nombres = {"valle_seco": "Valle, seco", "punta_seca": "Punta, seco",
               "punta_lluvia": "Punta + lluvia", "viernes_19h_lluvia_evento": "Viernes 19:00, lluvia + evento"}
    filas = []
    for r in mc.itertuples():
        filas.append([r.red, nombres[r.contexto], n(r.P_escasez_red, 3), n(r.precio_medio),
                      f"{n(r.p_eq_mediana)} [{n(r.p_eq_p5)} – {n(r.p_eq_p95)}]",
                      pct(r.P_precio_final_mayor_tarifa_alta)])
    inf.tabla("Monte Carlo Bayes → Euler (500 muestras por contexto)",
              ["Red", "Contexto", "P(escasez)", "Precio medio", "p* mediana [p5 – p95]",
               "P(p ≥ 13 Bs)"], filas)
    inf.parrafo(
        "La banda es ancha porque cada contexto deja abiertas varias combinaciones de demanda y "
        "oferta. Con la red calibrada el precio de equilibrio queda acotado en 14,14 Bs, que "
        "corresponde a demanda alta con oferta baja: tres niveles no alcanzan para representar "
        "el viernes con lluvia y evento, donde los datos muestran 21 Bs. Es una limitación de "
        "la discretización, no del modelo de Euler.")
    inf.figura("Banda de precio por contexto (Monte Carlo)", os.path.join(FIG, "montecarlo_bandas.png"))

    # ------------------------------------------------------------ 4
    inf.titulo("4.", "Datos")
    inf.parrafo(
        "El Escenario E se construyó con datos sintéticos generados con el modelo de Euler del "
        "propio caso (parámetros del Escenario A y efectos de contexto supuestos), porque no se "
        "encontraron series públicas de precios y solicitudes de plataformas de transporte en "
        "Cochabamba. Los únicos datos reales son los de precipitación de SENAMHI (diciembre 2023 "
        "y enero 2026) y las tarifas del transporte público municipal (Ley Municipal 1575/2024), "
        "que se usaron para acotar rangos. Los resultados obtenidos sobre estos datos validan el "
        "funcionamiento del sistema, no su capacidad predictiva en el mercado real.")
    filas = [["lluvia_senamhi.csv (fuente senamhi)", "Real", "Boletines de SENAMHI; parte leída de gráficos"],
             ["lluvia_senamhi.csv (supuesto)", "Supuesto", "Septiembre y octubre de 2026, sin documento"],
             ["cotizaciones_indrive.csv", "Sintético", "generar_sintetico.py, semilla 2026"],
             ["escenario_e.csv", "Derivado de sintético", "preparar.py"],
             ["Pasaje de 2,50 Bs y costo de 2,49 Bs", "Real (transporte público)",
              "Ley Municipal 1575/2024, Los Tiempos"]]
    inf.tabla("Origen de cada dato", ["Archivo o dato", "Estado", "Origen"], filas,
              fuente="Fuente: Elaboración propia en base a datos/README.md 2026")
    inf.parrafo(
        "Los efectos del contexto con que se generaron las cotizaciones (por ejemplo, hora punta "
        "+12 y lluvia +10 en la demanda base; lluvia −6 y tráfico alto −4 en la oferta base) son "
        "supuestos y están documentados en datos/README.md junto con el procedimiento para "
        "reemplazarlos por capturas reales. Para pasar a datos reales basta con cargar las "
        "nuevas cotizaciones con la columna fuente y volver a ejecutar preparar.py y los scripts "
        "de la sección 9.")

    # ------------------------------------------------------------ 5
    inf.titulo("5.", "Integración y pipeline")
    inf.parrafo(
        "main.py encadena las cuatro etapas de la Figura 1. Recibe el contexto por línea de "
        "comandos, obtiene D0, S0 y P(escasez) con contexto_a_parametros(), simula con "
        "simular(), estima el precio con el filtro de Kalman sobre mediciones ruidosas y aplica "
        "la regla de decisión. Imprime la decisión con su intervalo de credibilidad y guarda una "
        "figura. Ejemplos:")
    inf.vinetas(["python main.py --hora punta --clima lluvia --evento si  (Figura 9, "
                 "resultados/figuras/pipeline_punta_lluvia_si.png)",
                 "python main.py --hora valle --clima seco --evento no --red calibrada "
                 "--mediciones precio  (resultados/figuras/pipeline_valle_seco_no.png)",
                 "python main.py --escenario E --fecha 2026-10-30  (Figura 10, "
                 "resultados/figuras/pipeline_E_2026-10-30.png)"])
    inf.parrafo(
        "Con la red experta, el viernes a las 19:00 con lluvia y evento la red da P(escasez) = "
        f"{n(viernes_exp['P_escasez_red_experta'], 3)}, D0 = 62,0 y S0 = 45,4. El precio "
        "estimado ronda los 13 Bs, justo en el umbral, y la decisión alterna entre SUBIR e "
        "INCENTIVAR. En valle seco la decisión es MANTENER en todo el horizonte. En el modo "
        "--escenario E, a las 19:00 del viernes el precio estimado es 20,8 Bs con intervalo "
        "[20,0; 21,6], P(escasez) = 0,655 con la red calibrada y la decisión es INCENTIVAR.")
    inf.figura("Pipeline para el contexto viernes 19:00, lluvia y evento",
               os.path.join(FIG, "pipeline_punta_lluvia_si.png"))
    inf.figura("Pipeline sobre un día del Escenario E (datos sintéticos)",
               os.path.join(FIG, "pipeline_E_2026-10-30.png"), fuente=FUENTE_SINT)

    # ------------------------------------------------------------ 6
    inf.titulo("6.", "Decisión de tarifa (pregunta 7)")
    inf.titulo("6.1", "Regla", 2)
    inf.parrafo(
        f"En cada intervalo: si P(escasez) > X y el límite inferior del precio estimado supera "
        f"Y, se incentiva a los conductores; si P(escasez) > X y el límite inferior no supera Y, "
        f"se sube el precio; en otro caso se mantiene. Los valores elegidos son X = {n(X, 2)} e "
        f"Y = {n(Y, 0)} Bs. El límite inferior es la estimación de Kalman menos dos "
        "desviaciones. Se usa en lugar de la estimación puntual para que la afirmación «el "
        "precio ya supera Y» tenga 97,5 % de confianza: solo se renuncia a subir el precio, y "
        "se paga un bono, cuando es casi seguro que el precio ya es alto. Además, cada acción "
        f"se mantiene al menos {P.PERMANENCIA_MIN * 5} minutos antes de cambiar (sección 6.3). "
        "La regla está en decision.py.")

    inf.titulo("6.2", "Elección de X e Y", 2)
    inf.parrafo(
        "Se barrieron X de 0,05 a 0,95 e Y de 8 a 22 Bs sobre los escenarios B y C (con Kalman "
        "sobre mediciones simuladas) y E (siete días de datos sintéticos) "
        "(resultados/tablas/decision_barrido.csv). Cada acción se traduce en un cambio de D0 y "
        "S0 y se vuelve a simular con simular(): subir es un recargo de 2 Bs (menos demanda, más "
        "oferta) e incentivar es un bono de 2 Bs por viaje al conductor (más oferta).")
    inf.parrafo(
        "X se eligió con el índice de Youden (tasa de detección menos tasa de falsa alarma), "
        "donde los intervalos de estrés son la ventana de B y C y, en E, los intervalos con "
        f"tarifa alta (resultados/tablas/decision_eleccion_X.csv). El promedio forma una meseta "
        f"de {n(eX.loc[0.30, 'youden_promedio'], 3)} entre X = 0,25 y 0,35: por debajo la regla "
        "también actúa en valle seco y por encima de 0,65 ya no actúa ni en hora punta con "
        "lluvia. Se tomó el centro de la meseta.")
    inf.parrafo(
        "Y se eligió con el costo marginal de retener pasajeros "
        "(resultados/tablas/decision_eleccion_Y.csv): al bajar Y se incentiva en más intervalos "
        "y se sube en menos, de modo que se expulsan menos pasajeros pero se pagan más bonos. "
        f"El paso de 14 a 13 Bs tiene el menor costo, {n(eY.loc[13.0, 'costo_marginal_por_pasajero'], 1)} "
        "Bs de bono por pasajero retenido; el resto de la grilla está entre 65 y 93 Bs. La curva "
        "es plana entre 8 y 15 Bs, así que Y queda poco identificado y en la práctica es una "
        "decisión de cuánto se quiere proteger al pasajero; 13 Bs coincide además con el umbral "
        "de tarifa alta (1,3 p_ref). La Figura 11 resume el barrido "
        "(resultados/figuras/decision_barrido.png).")
    inf.figura("Barrido de los umbrales X e Y", os.path.join(FIG, "decision_barrido.png"),
               fuente=FUENTE_SINT)
    filas = []
    for e in "BCE":
        r = dec.loc[e]
        filas.append([e, pct(r.frac_MANTENER), pct(r.frac_SUBIR), pct(r.frac_INCENTIVAR),
                      pct(r.tasa_deteccion), pct(r.tasa_falsa_alarma),
                      n(r.delta_servicios_pct, 1) + " %",
                      f"{n(r.demanda_insatisfecha_base, 0)} → {n(r.demanda_insatisfecha, 0)}",
                      n(r.costo_incentivos_bs, 0)])
    inf.tabla(f"Resultado de la regla con X = {n(X, 2)} e Y = {n(Y, 0)} Bs",
              ["Esc.", "Mantener", "Subir", "Incentivar", "Detección", "Falsa alarma",
               "Δ servicios", "Demanda insatisfecha", "Bonos (Bs)"], filas, fuente=FUENTE_SINT)
    inf.parrafo(
        "La Tabla 11 (resultados/tablas/decision_resultado_elegido.csv) muestra que la regla "
        "detecta todos los intervalos de estrés de B y C y el 98 % de los de E, aumenta los "
        "servicios en B y C cerca de un 1 % y reduce la demanda insatisfecha en ambos. En los "
        "tres escenarios atrae pasajeros en lugar de expulsarlos, porque el bono aumenta la "
        "oferta, baja el precio de mercado y entra más demanda.")

    inf.titulo("6.3", "Permanencia mínima", 2)
    p1 = perm[(perm.permanencia == 1)].set_index("escenario")
    p6 = perm[(perm.permanencia == P.PERMANENCIA_MIN)].set_index("escenario")
    filas = [[int(r.permanencia) * 5, r.escenario, int(r.cambios_de_accion),
              n(r.delta_servicios_pct, 2) + " %",
              f"{n(r.demanda_insatisfecha_base, 0)} → {n(r.demanda_insatisfecha, 0)}"]
             for r in perm.itertuples()]
    inf.tabla("Efecto de la permanencia mínima de cada acción",
              ["Permanencia (min)", "Esc.", "Cambios de acción", "Δ servicios",
               "Demanda insatisfecha"], filas, fuente=FUENTE_SINT)
    inf.parrafo(
        "Sin permanencia mínima, la acción cambiaba "
        f"{int(p1.loc['C', 'cambios_de_accion'])} veces en 60 intervalos en C y "
        f"{int(p1.loc['E', 'cambios_de_accion'])} veces en E, porque el límite inferior del precio "
        "ronda los 13 Bs. Con la dinámica oscilante del modelo (r = −0,75), cada cambio produce "
        "un transitorio y la demanda insatisfecha en C subía de "
        f"{n(p1.loc['C', 'demanda_insatisfecha_base'], 0)} a {n(p1.loc['C', 'demanda_insatisfecha'], 0)}. "
        f"Con {P.PERMANENCIA_MIN * 5} minutos de permanencia baja a "
        f"{n(p6.loc['C', 'demanda_insatisfecha'], 0)} (Tabla 12, "
        "resultados/tablas/decision_permanencia.csv). En E la demanda insatisfecha todavía queda "
        f"un {pct(p6.loc['E', 'demanda_insatisfecha'] / p6.loc['E', 'demanda_insatisfecha_base'] - 1)} "
        "por encima de no actuar. La respuesta completa, con sus límites, está en "
        "resultados/pregunta7.md.")

    # ------------------------------------------------------------ 7
    inf.titulo("7.", "Limitaciones")
    inf.vinetas([
        "**Datos sintéticos.** El Escenario E valida que el sistema funciona de punta a punta, "
        "no que prediga el mercado real. Una semana, una ruta y eventos programados.",
        "**Parámetros no identificables con regresión simple.** Se usaron los valores nominales "
        "del documento; estimarlos con datos reales requeriría variables instrumentales o "
        "experimentos de precio.",
        "**Discretización en tres niveles.** La red no puede representar demandas más altas que "
        "la categoría alta (D0 = 65), por lo que subestima el precio en contextos extremos.",
        "**Escasez como D > S.** Casi no depende del contexto porque el precio equilibra el "
        "mercado; se usó tarifa_alta como proxy, que mide la presión sobre el precio y no los "
        "pasajeros sin servicio.",
        "**Evaluación de la decisión en lazo abierto.** Las decisiones se toman con las "
        "estimaciones de la trayectoria sin intervenir; en la realidad cada acción cambiaría las "
        "mediciones siguientes.",
        "**Recargo absorbido por el precio dinámico.** En el modelo lineal un recargo fijo baja "
        "el precio de mercado en la misma cantidad; solo el bono mueve el equilibrio.",
        "**Montos del bono y del recargo supuestos** (2 Bs cada uno, params.py).",
        "**Costos de operación no observados.** El nodo Costos de la red no se pudo calibrar.",
    ])

    # ------------------------------------------------------------ 8
    inf.titulo("8.", "Conclusiones")
    inf.vinetas([
        f"El modelo de Euler con los parámetros del documento es estable (Δt α (β + γ) = "
        f"{n(est['producto'])} < 2) pero oscilante (r = {n(est['r'])}); desde p0 = 10 converge a "
        f"p* = {n(met.loc['A', 'p_eq_teorico_base'])} Bs en {n(met.loc['A', 'tiempo_hasta_equilibrio_min'], 0)} "
        f"minutos. Los equilibrios de B y C son {n(met.loc['B', 'p_eq_teorico_ventana'])} y "
        f"{n(met.loc['C', 'p_eq_teorico_ventana'])} Bs.",
        f"Medir las solicitudes además del precio mejora la estimación de la demanda en "
        f"{pct(ganancia)} (Escenario D) y la predicción del precio en los datos de "
        f"{n(s.rmse_prediccion_1paso_precio)} a {n(c.rmse_prediccion_1paso_precio)} Bs (Escenario E).",
        "Integrar las capas mejora cada una: alimentar el filtro de Kalman con el D0 y S0 que "
        "entrega la red para cada contexto reduce el error de predicción del precio a menos de "
        "la mitad frente a usar valores fijos.",
        "En los datos la escasez D > S no distingue contextos; la tarifa alta sí, y es el "
        "indicador observable adecuado para calibrar la red.",
        f"La regla de decisión con X = {n(X, 2)}, Y = {n(Y, 0)} Bs y 30 minutos de permanencia "
        "detecta los intervalos de estrés y prefiere incentivar a los conductores cuando el "
        "precio ya es alto con seguridad. Pagar el bono cuesta unos 65 Bs por pasajero retenido, "
        "por lo que su justificación es proteger al pasajero de precios altos más que la "
        "eficiencia.",
    ])

    # ------------------------------------------------------------ 9
    inf.titulo("9.", "Reproducibilidad")
    inf.parrafo(
        "Todo el trabajo se reproduce desde cero con Python 3.11 o superior. Los resultados son "
        f"deterministas (semilla {P.SEMILLA}):")
    inf.vinetas(["pip install -r requirements.txt",
                 "python escenarios.py", "python escenario_e.py", "python montecarlo.py",
                 "python decision.py", "python main.py", "python main.py --escenario E",
                 "python informe/generar_informe.py"], estilo="List Number")
    inf.parrafo(
        "La Tabla 13 lista los archivos de resultados que se citan en este informe y el script "
        "que genera cada uno.")
    origen = {"metricas_escenarios.csv": "escenarios.py", "kalman_valor_informacion.csv": "escenarios.py",
              "estimaciones_kalman_D.csv": "escenarios.py", "escenarios_ABC.png": "escenarios.py",
              "escenario_A_convergencia.png": "escenarios.py", "escenario_D_kalman.png": "escenarios.py",
              "pregunta7.md": "redactado a partir de decision.py"}
    filas = []
    for sub in ("figuras", "tablas"):
        for f in sorted(os.listdir(os.path.join(P.RESULTADOS, sub))):
            if f.startswith("E_"):
                gen = "escenario_e.py"
            elif f.startswith("montecarlo"):
                gen = "montecarlo.py"
            elif f.startswith("decision"):
                gen = "decision.py"
            elif f.startswith("pipeline"):
                gen = "main.py"
            else:
                gen = origen.get(f, "escenarios.py")
            filas.append([f"resultados/{sub}/{f}", gen])
    filas.append(["resultados/pregunta7.md", origen["pregunta7.md"]])
    inf.tabla("Archivos de resultados", ["Archivo", "Generado por"], filas)

    # ------------------------------------------------------------ Referencias
    inf.titulo("", "Referencias")
    for ref in [
        "Documento del caso de estudio «Precio dinámico de transporte según oferta y demanda». "
        "Asignatura Modelado, 2026. [completar autor]",
        "Los Tiempos (2024-2025). Notas sobre la Ley Municipal 1575/2024 de tarifas del "
        "transporte público en Cochabamba (pasaje de 2,50 Bs y estudio de costos de 2,49 Bs).",
        "SENAMHI (2024). Boletín BOL-PC-N01-2024 (diciembre de 2023): Tablas 1 a 4 e "
        "Ilustración 5. Servicio Nacional de Meteorología e Hidrología de Bolivia.",
        "SENAMHI (2026). Boletín BOL-RC-N01-2026 (enero de 2026): Gráficos 16 a 18 y 37. "
        "Servicio Nacional de Meteorología e Hidrología de Bolivia.",
    ]:
        inf.parrafo(ref)


# ---------------------------------------------------------------------------
# Presentacion
# ---------------------------------------------------------------------------
def presentacion():
    from pptx import Presentation
    from pptx.util import Inches, Pt as PPt
    from pptx.dml.color import RGBColor as PRGB

    met = tabla_csv("metricas_escenarios.csv").set_index("escenario")
    dec = tabla_csv("decision_resultado_elegido.csv").set_index("escenario")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def diapositiva(titulo, puntos=None, imagen=None):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = titulo
        for p in s.shapes.title.text_frame.paragraphs:
            p.font.size = PPt(30)
            p.font.name = "Times New Roman"
            p.font.color.rgb = PRGB(0, 0, 0)
        izq = Inches(0.6)
        if puntos:
            ancho = Inches(5.6) if imagen else Inches(12)
            tb = s.shapes.add_textbox(izq, Inches(1.5), ancho, Inches(5.5)).text_frame
            tb.word_wrap = True
            for i, pt in enumerate(puntos):
                par = tb.paragraphs[0] if i == 0 else tb.add_paragraph()
                par.text = "• " + pt
                par.font.size = PPt(18)
                par.font.name = "Times New Roman"
                par.space_after = PPt(8)
        if imagen:
            x = Inches(6.4) if puntos else Inches(1.2)
            w = Inches(6.6) if puntos else Inches(10.9)
            s.shapes.add_picture(imagen, x, Inches(1.5), width=w)
        return s

    diapositiva("Precio dinámico de transporte",
                ["Euler + red bayesiana + filtro de Kalman + regla de decisión",
                 "Asignatura Modelado, 2026",
                 "Persona 4: escenarios, datos, decisión e integración"],
                os.path.join(FIG_INF, "diagrama_pipeline.png"))
    diapositiva("Modelo de Euler",
                [f"p* = {n(met.loc['A', 'p_eq_teorico_base'])} Bs; Δt α (β + γ) = 1,75 < 2: estable",
                 "r = −0,75: converge oscilando desde p0 = 10",
                 f"Equilibrio en {n(met.loc['A', 'tiempo_hasta_equilibrio_min'], 0)} min"],
                os.path.join(FIG, "escenario_A_convergencia.png"))
    diapositiva("Escenarios B y C",
                [f"B (D0 = 100): p* = {n(met.loc['B', 'p_eq_teorico_ventana'])} Bs, máx. "
                 f"{n(met.loc['B', 'precio_max'])} Bs",
                 f"C (S0 = 30): p* = {n(met.loc['C', 'p_eq_teorico_ventana'])} Bs, máx. "
                 f"{n(met.loc['C', 'precio_max'])} Bs",
                 "Perturbación entre los minutos 60 y 180"],
                os.path.join(FIG, "escenarios_ABC.png"))
    diapositiva("Red bayesiana",
                ["Contexto → Tráfico → Demanda / Oferta → Escasez",
                 "Viernes 19:00 con lluvia y evento: P(escasez) = 0,69 (red experta)",
                 "Tabla de traducción calibrada: D0 = 47/54/65, S0 = 43/47/50"],
                os.path.join(FIG_INF, "diagrama_red_bayesiana.png"))
    diapositiva("Escenario D: filtro de Kalman",
                ["Solo precio frente a precio + solicitudes",
                 "Las solicitudes reducen el error de la demanda casi a la mitad",
                 "Cobertura de ±2 desviaciones ≈ 95 %"],
                os.path.join(FIG, "escenario_D_kalman.png"))
    diapositiva("Escenario E (datos sintéticos)",
                ["La escasez D > S no depende del contexto: el precio equilibra",
                 "La tarifa alta sí: 1 % en valle, 63 % en punta",
                 "Regresión simple no recupera α, β, γ (identificabilidad)"],
                os.path.join(FIG, "E_frecuencias_contexto.png"))
    diapositiva("Kalman sobre los datos",
                ["D0 y S0 del contexto (red calibrada) en el término B u",
                 "Error de predicción a un paso: 1,47 Bs; NIS ≈ 2",
                 "Solo precio: 2,63 Bs"],
                os.path.join(FIG, "E_kalman_viernes.png"))
    diapositiva("Monte Carlo Bayes → Euler", None, os.path.join(FIG, "montecarlo_bandas.png"))
    diapositiva("Decisión de tarifa (pregunta 7)",
                [f"INCENTIVAR si P(escasez) > {n(P.UMBRAL_X, 2)} y el límite inferior del precio > "
                 f"{n(P.UMBRAL_Y, 0)} Bs",
                 "SUBIR si P(escasez) > X y el precio aún no es alto con seguridad",
                 f"B: +{n(dec.loc['B', 'delta_servicios_pct'], 1)} % servicios; C: "
                 f"+{n(dec.loc['C', 'delta_servicios_pct'], 1)} %",
                 "Permanencia mínima de 30 min para evitar parpadeo"],
                os.path.join(FIG, "decision_barrido.png"))
    diapositiva("Conclusiones y limitaciones",
                ["Integrar las capas mejora cada una (Kalman con contexto: error a menos de la mitad)",
                 "La tarifa alta es el indicador observable de escasez",
                 "El bono cuesta ~65 Bs por pasajero retenido: protege al pasajero, no es eficiente",
                 "Datos sintéticos: validan el sistema, no predicen el mercado real",
                 "Reproducible: pip install -r requirements.txt y los scripts del README"])
    ruta = os.path.join(CARPETA, "presentacion.pptx")
    prs.save(ruta)
    return ruta


def comprobar_referencias(inf):
    texto = "\n".join(inf.texto_total)
    faltan = []
    for sub in ("figuras", "tablas"):
        for f in os.listdir(os.path.join(P.RESULTADOS, sub)):
            if f not in texto:
                faltan.append(f"resultados/{sub}/{f}")
    if "pregunta7.md" not in texto:
        faltan.append("resultados/pregunta7.md")
    return faltan


def main():
    os.makedirs(FIG_INF, exist_ok=True)
    inf = Informe()
    construir(inf)
    ruta = os.path.join(CARPETA, "informe.docx")
    inf.doc.save(ruta)
    print(f"Informe: {ruta}  ({inf.n_tabla} tablas, {inf.n_figura} figuras)")
    faltan = comprobar_referencias(inf)
    if faltan:
        raise SystemExit(f"Archivos de resultados/ no citados en el informe: {faltan}")
    print("Todos los archivos de resultados/ estan citados en el informe.")
    print(f"Presentacion: {presentacion()}")


if __name__ == "__main__":
    main()
